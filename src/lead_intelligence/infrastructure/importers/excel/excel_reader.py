"""The Excel adapter: turns an .xlsx workbook into an ImportedLeadDataset.

WHY THIS FILE EXISTS, AND WHY IT USES openpyxl DIRECTLY (NOT pandas):
Data-profiling of the reference file found that pandas.read_excel silently
destroyed leading zeros in the D-U-N-S identifier column (turning the true
9-digit text value "036611498" into the integer 36611498) — because pandas
performs its own column-wide type inference on top of what's in the file.
openpyxl, read cell-by-cell, returns exactly the type Excel itself recorded
for each cell (a string cell stays a string, leading zeros and all), with
no extra inference layer. So the rule this class follows is simple and
general, not specific to any one column: *never coerce a cell's value —
always use whatever native type openpyxl already determined.* That one
rule is what "preserve identifiers as strings" and "preserve leading
zeros" actually mean in practice, for this column or any other.

This class's responsibility stops at faithfully reporting what's in the
file. It does not decide whether a value is valid, does not rename
columns, does not drop or merge rows, and does not guess at missing data.
"""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl
from loguru import logger

from lead_intelligence.application.dto.models import (
    ImportedLeadDataset,
    ImportWarning,
    RawRecord,
    SourceMetadata,
)
from lead_intelligence.application.ports.source_reader_port import SourceReaderPort
from lead_intelligence.domain.exceptions import CorruptedSourceError, EmptySourceError
from lead_intelligence.infrastructure.importers.excel.file_validator import (
    ExcelFileValidator,
)
from lead_intelligence.infrastructure.importers.excel.sheet_selector import (
    ExcelSheetSelector,
    SheetSummary,
)


class ExcelSourceReader(SourceReaderPort):
    """Reads lead/contact data out of a single .xlsx workbook.

    Composes an ExcelFileValidator (structural file checks) and an
    ExcelSheetSelector (sheet-choice logic) as collaborators; this class's
    own job is solely the openpyxl I/O and turning worksheet rows into
    RawRecord objects.
    """

    def __init__(
        self,
        file_path: str | Path,
        *,
        sheet_name: str | None = None,
        file_validator: ExcelFileValidator | None = None,
        sheet_selector: ExcelSheetSelector | None = None,
    ) -> None:
        """Configure a reader bound to one Excel file.

        Args:
            file_path: Path to the .xlsx file to read.
            sheet_name: Optional sheet to read every time `read()` is
                called without an explicit sheet_name argument. If omitted,
                the sheet is selected automatically on each read().
            file_validator: Injectable validator, for testing. Defaults to
                a standard ExcelFileValidator.
            sheet_selector: Injectable selector, for testing. Defaults to
                a standard ExcelSheetSelector.
        """

        self._file_path = Path(file_path)
        self._preferred_sheet = sheet_name
        self._file_validator = file_validator or ExcelFileValidator()
        self._sheet_selector = sheet_selector or ExcelSheetSelector()

    def list_available_sheets(self) -> list[str]:
        """Return the names of every sheet in the workbook.

        Raises:
            SourceNotFoundError, UnsupportedSourceFormatError,
            CorruptedSourceError: If the file cannot be opened.
        """

        self._file_validator.validate(self._file_path)
        workbook = self._open_workbook()
        try:
            return list(workbook.sheetnames)
        finally:
            workbook.close()

    def read(self, sheet_name: str | None = None) -> ImportedLeadDataset:
        """Read one worksheet and return it as an ImportedLeadDataset.

        Args:
            sheet_name: Sheet to read. If omitted, falls back to the
                sheet_name given at construction time, or automatic
                selection if neither is set.

        Returns:
            An ImportedLeadDataset with every row as an unmodified
            RawRecord, source metadata, and any structural warnings.

        Raises:
            SourceNotFoundError, UnsupportedSourceFormatError,
            CorruptedSourceError: If the file cannot be opened.
            SheetSelectionError: If no single sheet can be determined.
            EmptySourceError: If the selected sheet has no data rows.
        """

        logger.info("Reading Excel source: {}", self._file_path)
        self._file_validator.validate(self._file_path)
        workbook = self._open_workbook()
        try:
            summaries = [
                self._summarize(workbook[name]) for name in workbook.sheetnames
            ]
            effective_preference = sheet_name or self._preferred_sheet
            selected_name, selection_warnings = self._sheet_selector.select(
                summaries, effective_preference
            )

            worksheet = workbook[selected_name]
            header, records, row_warnings = self._read_rows(worksheet, selected_name)

            if not records:
                raise EmptySourceError(
                    f"Sheet '{selected_name}' in {self._file_path} contains no data rows."
                )

            metadata = SourceMetadata(
                source_path=str(self._file_path),
                source_type="excel",
                sheet_name=selected_name,
                available_sheets=tuple(workbook.sheetnames),
                column_names=tuple(header),
                row_count=len(records),
                imported_at=datetime.now(timezone.utc),
            )
            warnings = tuple(selection_warnings + row_warnings)

            logger.info(
                "Imported {} record(s) from sheet '{}' of {} ({} warning(s))",
                len(records),
                selected_name,
                self._file_path.name,
                len(warnings),
            )
            return ImportedLeadDataset(
                records=tuple(records), metadata=metadata, warnings=warnings
            )
        finally:
            workbook.close()

    def _open_workbook(self):
        """Open the workbook read-only, translating library errors on the way.

        read_only mode streams rows instead of loading the whole workbook
        into memory, which matters for files with thousands of rows.
        data_only=True returns formulas' last-calculated *values* rather
        than the formula text, since we want the data a user would see.
        """

        try:
            return openpyxl.load_workbook(
                self._file_path, read_only=True, data_only=True
            )
        except Exception as exc:  # noqa: BLE001 - translate to domain vocabulary
            raise CorruptedSourceError(
                f"Could not open '{self._file_path}' as a valid Excel workbook: {exc}"
            ) from exc

    @staticmethod
    def _summarize(worksheet: Any) -> SheetSummary:
        """Compute cheap row/column counts for one worksheet, for selection."""

        return SheetSummary(
            name=worksheet.title,
            row_count=max((worksheet.max_row or 0) - 1, 0),
            column_count=worksheet.max_column or 0,
        )

    def _read_rows(
        self, worksheet: Any, sheet_name: str
    ) -> tuple[list[str], list[RawRecord], list[ImportWarning]]:
        """Read a worksheet's header and data rows into RawRecords.

        Every column from the header row is preserved exactly as written
        (see _build_header). A data row where every cell is blank is
        skipped and reported as a warning — this handles Excel's tendency
        to report a "used range" larger than the real data — but no
        populated cell, row, or column is ever altered or dropped.
        """

        row_iterator = worksheet.iter_rows(values_only=True)
        try:
            raw_header = next(row_iterator)
        except StopIteration:
            return [], [], []

        header, warnings = self._build_header(raw_header)

        records: list[RawRecord] = []
        for row_number, row_values in enumerate(row_iterator, start=2):
            if all(value is None for value in row_values):
                warnings.append(
                    ImportWarning(
                        code="SKIPPED_BLANK_ROW",
                        message=f"Row {row_number} is entirely blank and was skipped.",
                        row_number=row_number,
                    )
                )
                continue

            values = dict(zip(header, row_values))
            records.append(
                RawRecord(row_number=row_number, sheet_name=sheet_name, values=values)
            )

        return header, records, warnings

    @staticmethod
    def _build_header(
        raw_header: tuple[Any, ...]
    ) -> tuple[list[str], list[ImportWarning]]:
        """Turn a raw header row into column names, unmodified except where blank.

        Column names are preserved literally (no stripping, no casing
        changes). A blank header cell gets a positional placeholder
        ("Unnamed: N") purely so it can serve as a dict key — this is
        flagged with a warning, never done silently. A duplicate header
        name is also flagged: a plain dict cannot hold two values under one
        key, so the *last* occurrence's value wins in RawRecord.values,
        while the full original header list (duplicates included) still
        lives, unmodified, in SourceMetadata.column_names.
        """

        header: list[str] = []
        warnings: list[ImportWarning] = []
        first_seen_at: dict[str, int] = {}

        for index, raw_name in enumerate(raw_header, start=1):
            if raw_name is None:
                column_name = f"Unnamed: {index}"
                warnings.append(
                    ImportWarning(
                        code="BLANK_COLUMN_HEADER",
                        message=(
                            f"Column {index} has no header text; using placeholder "
                            f"'{column_name}'."
                        ),
                        column_name=column_name,
                    )
                )
            else:
                column_name = str(raw_name)

            if column_name in first_seen_at:
                warnings.append(
                    ImportWarning(
                        code="DUPLICATE_COLUMN_HEADER",
                        message=(
                            f"Column header '{column_name}' appears more than once "
                            f"(columns {first_seen_at[column_name]} and {index}); "
                            "only the last occurrence's value is retained per row."
                        ),
                        column_name=column_name,
                    )
                )
            else:
                first_seen_at[column_name] = index

            header.append(column_name)

        return header, warnings
